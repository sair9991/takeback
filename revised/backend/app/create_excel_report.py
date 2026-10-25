#!/usr/bin/env python3
"""
Create comprehensive, professionally formatted Excel workbook for:
1. Overall Model Benchmarks (Move accuracy + Decision Time predictions)
2. Snapshot Scaling Evaluation (Convergence across varying snapshot counts)
3. Individual Player Scaling Sheets (saiteja1000, Schwoby_1200, hikaru, stuffyness)
"""

import csv
from pathlib import Path
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

BASE_DIR = Path(__file__).resolve().parent
BENCHMARK_CSV = BASE_DIR / "model_benchmark_accuracy.csv"
SCALING_CSV = BASE_DIR / "snapshot_scaling_evaluation.csv"
OUTPUT_XLSX = BASE_DIR / "chess_clone_model_evaluations.xlsx"

# Color Palette (Corporate Navy & Slate)
NAVY_HEADER = "1E3A8A"       # Primary Dark Navy
STEEL_HEADER = "0F766E"      # Deep Teal / Steel for secondary
SECTION_HEADER = "334155"    # Slate Gray for subheadings
ACCENT_GREEN = "DCFCE7"      # Soft Green for positive improvements
ACCENT_ROSE = "FFE4E6"       # Soft Rose for slight dips
ZEBRA_FILL = "F8FAFC"        # Off-white / light slate zebra striping
WHITE_FILL = "FFFFFF"
BORDER_COLOR = "CBD5E1"      # Subtle light gray border

font_title = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
font_subtitle = Font(name="Calibri", size=11, italic=True, color="E2E8F0")
font_section = Font(name="Calibri", size=13, bold=True, color="1E293B")
font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
font_bold = Font(name="Calibri", size=11, bold=True)
font_regular = Font(name="Calibri", size=11)
font_positive = Font(name="Calibri", size=11, bold=True, color="166534")
font_negative = Font(name="Calibri", size=11, color="991B1B")

fill_title = PatternFill(start_color=NAVY_HEADER, end_color=NAVY_HEADER, fill_type="solid")
fill_header = PatternFill(start_color=NAVY_HEADER, end_color=NAVY_HEADER, fill_type="solid")
fill_header_teal = PatternFill(start_color=STEEL_HEADER, end_color=STEEL_HEADER, fill_type="solid")
fill_zebra = PatternFill(start_color=ZEBRA_FILL, end_color=ZEBRA_FILL, fill_type="solid")
fill_white = PatternFill(start_color=WHITE_FILL, end_color=WHITE_FILL, fill_type="solid")
fill_highlight = PatternFill(start_color=ACCENT_GREEN, end_color=ACCENT_GREEN, fill_type="solid")

thin_side = Side(border_style="thin", color=BORDER_COLOR)
thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
double_bottom_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=Side(border_style="double", color="64748B"))

align_center = Alignment(horizontal="center", vertical="center")
align_left = Alignment(horizontal="left", vertical="center")
align_right = Alignment(horizontal="right", vertical="center")
align_header = Alignment(horizontal="center", vertical="center", wrap_text=True)


def auto_fit_columns(ws, min_width=12, max_width=40):
    ws.views.sheetView[0].showGridLines = True
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            val_str = str(cell.value or "")
            if cell.number_format and "%" in cell.number_format:
                val_str += "%"
            max_len = max(max_len, len(val_str))
        ws.column_dimensions[col_letter].width = min(max(max_len + 4, min_width), max_width)


def create_overall_benchmarks_sheet(wb, bench_rows):
    ws = wb.active
    ws.title = "Overall Model Benchmarks"

    # Title Banner
    ws.merge_cells("A1:L1")
    title_cell = ws["A1"]
    title_cell.value = "CHESS CLONE NEURAL NETWORK BENCHMARK EVALUATIONS"
    title_cell.font = font_title
    title_cell.fill = fill_title
    title_cell.alignment = align_center
    ws.row_dimensions[1].height = 35

    ws.merge_cells("A2:L2")
    sub_cell = ws["A2"]
    sub_cell.value = "Comparative Benchmark: Baseline Maia-3 vs. Personalized Clone Models (25 Games Evaluated per Player)"
    sub_cell.font = font_subtitle
    sub_cell.fill = fill_title
    sub_cell.alignment = align_center
    ws.row_dimensions[2].height = 22

    # --- Table 1: Move Prediction ---
    start_r = 4
    ws.cell(row=start_r, column=1, value="1. Move Prediction Performance (Baseline Maia-3 vs. Clone Model)").font = font_section
    ws.row_dimensions[start_r].height = 25

    t1_headers = [
        "Player", "Analyzed Games", "Evaluated Positions", "Maia Top-1", "Clone Top-1",
        "Top-1 Δ", "Maia Top-3", "Clone Top-3", "Top-3 Δ", "Maia Loss", "Clone Loss", "Loss Δ"
    ]
    header_r = start_r + 1
    ws.row_dimensions[header_r].height = 28
    for c_idx, h in enumerate(t1_headers, start=1):
        cell = ws.cell(row=header_r, column=c_idx, value=h)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_header
        cell.border = thin_border

    curr_r = header_r + 1
    for r_idx, r in enumerate(bench_rows):
        fill = fill_zebra if r_idx % 2 == 1 else fill_white
        ws.row_dimensions[curr_r].height = 22

        c1 = ws.cell(row=curr_r, column=1, value=r["username"])
        c1.alignment = align_center
        c1.font = font_bold
        c1.fill = fill
        c1.border = thin_border

        c2 = ws.cell(row=curr_r, column=2, value=int(r["games_analyzed"]))
        c2.number_format = "#,##0"
        c2.alignment = align_center
        c2.fill = fill
        c2.border = thin_border

        c3 = ws.cell(row=curr_r, column=3, value=int(r["positions_evaluated"]))
        c3.number_format = "#,##0"
        c3.alignment = align_center
        c3.fill = fill
        c3.border = thin_border

        # Maia Top 1
        c4 = ws.cell(row=curr_r, column=4, value=float(r["maia_top1_pct"]) / 100.0)
        c4.number_format = "0.00%"
        c4.alignment = align_right
        c4.fill = fill
        c4.border = thin_border

        # Clone Top 1
        c5 = ws.cell(row=curr_r, column=5, value=float(r["clone_top1_pct"]) / 100.0)
        c5.number_format = "0.00%"
        c5.alignment = align_right
        c5.font = font_bold
        c5.fill = fill
        c5.border = thin_border

        # Top 1 Delta
        d1 = float(r["top1_delta_pct"]) / 100.0
        c6 = ws.cell(row=curr_r, column=6, value=d1)
        c6.number_format = "+0.00%;-0.00%;0.00%"
        c6.alignment = align_right
        c6.font = font_positive if d1 > 0 else (font_negative if d1 < 0 else font_regular)
        c6.fill = fill_highlight if d1 > 0 else fill
        c6.border = thin_border

        # Maia Top 3
        c7 = ws.cell(row=curr_r, column=7, value=float(r["maia_top3_pct"]) / 100.0)
        c7.number_format = "0.00%"
        c7.alignment = align_right
        c7.fill = fill
        c7.border = thin_border

        # Clone Top 3
        c8 = ws.cell(row=curr_r, column=8, value=float(r["clone_top3_pct"]) / 100.0)
        c8.number_format = "0.00%"
        c8.alignment = align_right
        c8.font = font_bold
        c8.fill = fill
        c8.border = thin_border

        # Top 3 Delta
        d3 = float(r["top3_delta_pct"]) / 100.0
        c9 = ws.cell(row=curr_r, column=9, value=d3)
        c9.number_format = "+0.00%;-0.00%;0.00%"
        c9.alignment = align_right
        c9.font = font_positive if d3 > 0 else (font_negative if d3 < 0 else font_regular)
        c9.fill = fill_highlight if d3 > 0 else fill
        c9.border = thin_border

        # Maia Loss
        c10 = ws.cell(row=curr_r, column=10, value=float(r["maia_loss"]))
        c10.number_format = "0.0000"
        c10.alignment = align_right
        c10.fill = fill
        c10.border = thin_border

        # Clone Loss
        c11 = ws.cell(row=curr_r, column=11, value=float(r["clone_loss"]))
        c11.number_format = "0.0000"
        c11.alignment = align_right
        c11.font = font_bold
        c11.fill = fill
        c11.border = thin_border

        # Loss Delta
        ld = float(r["loss_reduction"])
        c12 = ws.cell(row=curr_r, column=12, value=-ld)
        c12.number_format = "-0.0000;+0.0000;0.0000"
        c12.alignment = align_right
        c12.font = font_positive if ld > 0 else (font_negative if ld < 0 else font_regular)
        c12.fill = fill_highlight if ld > 0 else fill
        c12.border = thin_border

        curr_r += 1

    # --- Table 2: Decision Time Prediction ---
    start_r2 = curr_r + 2
    ws.cell(row=start_r2, column=1, value="2. Decision Time Prediction Accuracy (Log-Normal Tempo Head vs. Game Clocks)").font = font_section
    ws.row_dimensions[start_r2].height = 25

    t2_headers = [
        "Player", "Evaluated Clocks", "Actual Median", "Pred Median", "Tempo Δ",
        "MedAE", "MAE", "RMSE", "Within ±1s", "Within ±2s", "Within ±5s", "Pearson r"
    ]
    header_r2 = start_r2 + 1
    ws.row_dimensions[header_r2].height = 28
    for c_idx, h in enumerate(t2_headers, start=1):
        cell = ws.cell(row=header_r2, column=c_idx, value=h)
        cell.font = font_header
        cell.fill = fill_header_teal
        cell.alignment = align_header
        cell.border = thin_border

    curr_r2 = header_r2 + 1
    for r_idx, r in enumerate(bench_rows):
        fill = fill_zebra if r_idx % 2 == 1 else fill_white
        ws.row_dimensions[curr_r2].height = 22

        c1 = ws.cell(row=curr_r2, column=1, value=r["username"])
        c1.alignment = align_center
        c1.font = font_bold
        c1.fill = fill
        c1.border = thin_border

        c2 = ws.cell(row=curr_r2, column=2, value=int(r["valid_time_positions"]))
        c2.number_format = "#,##0"
        c2.alignment = align_center
        c2.fill = fill
        c2.border = thin_border

        c3 = ws.cell(row=curr_r2, column=3, value=float(r["actual_median_time_s"]))
        c3.number_format = '0.00"s"'
        c3.alignment = align_right
        c3.fill = fill
        c3.border = thin_border

        c4 = ws.cell(row=curr_r2, column=4, value=float(r["pred_median_time_s"]))
        c4.number_format = '0.00"s"'
        c4.alignment = align_right
        c4.font = font_bold
        c4.fill = fill
        c4.border = thin_border

        c5 = ws.cell(row=curr_r2, column=5, value=float(r["median_tempo_delta_s"]))
        c5.number_format = '+0.00"s";-0.00"s";0.00"s"'
        c5.alignment = align_right
        c5.fill = fill
        c5.border = thin_border

        c6 = ws.cell(row=curr_r2, column=6, value=float(r["medae_s"]))
        c6.number_format = '0.00"s"'
        c6.alignment = align_right
        c6.fill = fill
        c6.border = thin_border

        c7 = ws.cell(row=curr_r2, column=7, value=float(r["mae_s"]))
        c7.number_format = '0.00"s"'
        c7.alignment = align_right
        c7.fill = fill
        c7.border = thin_border

        c8 = ws.cell(row=curr_r2, column=8, value=float(r["rmse_s"]))
        c8.number_format = '0.00"s"'
        c8.alignment = align_right
        c8.fill = fill
        c8.border = thin_border

        c9 = ws.cell(row=curr_r2, column=9, value=float(r["within_1s_pct"]) / 100.0)
        c9.number_format = "0.0%"
        c9.alignment = align_right
        c9.fill = fill
        c9.border = thin_border

        c10 = ws.cell(row=curr_r2, column=10, value=float(r["within_2s_pct"]) / 100.0)
        c10.number_format = "0.0%"
        c10.alignment = align_right
        c10.fill = fill
        c10.border = thin_border

        c11 = ws.cell(row=curr_r2, column=11, value=float(r["within_5s_pct"]) / 100.0)
        c11.number_format = "0.0%"
        c11.alignment = align_right
        c11.fill = fill
        c11.border = thin_border

        corr = float(r["time_correlation_r"])
        c12 = ws.cell(row=curr_r2, column=12, value=corr)
        c12.number_format = "0.000"
        c12.alignment = align_right
        c12.font = font_bold if corr >= 0.4 else font_regular
        c12.fill = fill_highlight if corr >= 0.4 else fill
        c12.border = thin_border

        curr_r2 += 1

    auto_fit_columns(ws, min_width=13)


def create_scaling_master_sheet(wb, scale_rows):
    ws = wb.create_sheet(title="Snapshot Scaling (All Users)")

    # Title Banner
    ws.merge_cells("A1:N1")
    title_cell = ws["A1"]
    title_cell.value = "SNAPSHOT SCALING & SAMPLE SIZE CONVERGENCE EVALUATION"
    title_cell.font = font_title
    title_cell.fill = fill_title
    title_cell.alignment = align_center
    ws.row_dimensions[1].height = 32

    headers = [
        "User", "Snapshots Evaluated", "Total Positions", "Maia Top-1", "Clone Top-1", "Top-1 Δ",
        "Maia Top-3", "Clone Top-3", "Maia Loss", "Clone Loss", "Loss Δ",
        "Pred Tempo", "Actual Tempo", "MedAE", "Within ±2s", "Pearson r"
    ]

    ws.row_dimensions[3].height = 28
    for c_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=3, column=c_idx, value=h)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_header
        cell.border = thin_border

    curr_r = 4
    for r_idx, r in enumerate(scale_rows):
        fill = fill_zebra if r_idx % 2 == 1 else fill_white
        ws.row_dimensions[curr_r].height = 20

        # Username
        c1 = ws.cell(row=curr_r, column=1, value=r["display_username"])
        c1.alignment = align_center
        c1.font = font_bold
        c1.fill = fill
        c1.border = thin_border

        # Snapshots
        c2 = ws.cell(row=curr_r, column=2, value=int(r["snapshots_evaluated"]))
        c2.number_format = "#,##0"
        c2.alignment = align_center
        c2.font = font_bold
        c2.fill = fill
        c2.border = thin_border

        # Total positions
        c3 = ws.cell(row=curr_r, column=3, value=int(r["total_available_positions"]))
        c3.number_format = "#,##0"
        c3.alignment = align_center
        c3.fill = fill
        c3.border = thin_border

        # Maia Top 1
        c4 = ws.cell(row=curr_r, column=4, value=float(r["maia_top1_pct"]) / 100.0)
        c4.number_format = "0.00%"
        c4.alignment = align_right
        c4.fill = fill
        c4.border = thin_border

        # Clone Top 1
        c5 = ws.cell(row=curr_r, column=5, value=float(r["clone_top1_pct"]) / 100.0)
        c5.number_format = "0.00%"
        c5.alignment = align_right
        c5.font = font_bold
        c5.fill = fill
        c5.border = thin_border

        # Top 1 Delta
        d1 = float(r["top1_delta_pct"]) / 100.0
        c6 = ws.cell(row=curr_r, column=6, value=d1)
        c6.number_format = "+0.00%;-0.00%;0.00%"
        c6.alignment = align_right
        c6.font = font_positive if d1 > 0 else (font_negative if d1 < 0 else font_regular)
        c6.fill = fill_highlight if d1 > 0 else fill
        c6.border = thin_border

        # Maia Top 3
        c7 = ws.cell(row=curr_r, column=7, value=float(r["maia_top3_pct"]) / 100.0)
        c7.number_format = "0.00%"
        c7.alignment = align_right
        c7.fill = fill
        c7.border = thin_border

        # Clone Top 3
        c8 = ws.cell(row=curr_r, column=8, value=float(r["clone_top3_pct"]) / 100.0)
        c8.number_format = "0.00%"
        c8.alignment = align_right
        c8.font = font_bold
        c8.fill = fill
        c8.border = thin_border

        # Maia Loss
        c9 = ws.cell(row=curr_r, column=9, value=float(r["maia_loss"]))
        c9.number_format = "0.0000"
        c9.alignment = align_right
        c9.fill = fill
        c9.border = thin_border

        # Clone Loss
        c10 = ws.cell(row=curr_r, column=10, value=float(r["clone_loss"]))
        c10.number_format = "0.0000"
        c10.alignment = align_right
        c10.font = font_bold
        c10.fill = fill
        c10.border = thin_border

        # Loss Delta
        ld = float(r["loss_reduction"])
        c11 = ws.cell(row=curr_r, column=11, value=-ld)
        c11.number_format = "-0.0000;+0.0000;0.0000"
        c11.alignment = align_right
        c11.font = font_positive if ld > 0 else (font_negative if ld < 0 else font_regular)
        c11.fill = fill_highlight if ld > 0 else fill
        c11.border = thin_border

        # Pred Tempo
        c12 = ws.cell(row=curr_r, column=12, value=float(r["pred_median_sec"]))
        c12.number_format = '0.00"s"'
        c12.alignment = align_right
        c12.fill = fill
        c12.border = thin_border

        # Actual Tempo
        c13 = ws.cell(row=curr_r, column=13, value=float(r["actual_median_sec"]))
        c13.number_format = '0.00"s"'
        c13.alignment = align_right
        c13.fill = fill
        c13.border = thin_border

        # MedAE
        c14 = ws.cell(row=curr_r, column=14, value=float(r["medae_sec"]))
        c14.number_format = '0.00"s"'
        c14.alignment = align_right
        c14.fill = fill
        c14.border = thin_border

        # Within 2s
        c15 = ws.cell(row=curr_r, column=15, value=float(r["within_2s_pct"]) / 100.0)
        c15.number_format = "0.0%"
        c15.alignment = align_right
        c15.fill = fill
        c15.border = thin_border

        # Pearson r
        corr = float(r["correlation_r"])
        c16 = ws.cell(row=curr_r, column=16, value=corr)
        c16.number_format = "0.000"
        c16.alignment = align_right
        c16.font = font_bold if corr >= 0.4 else font_regular
        c16.fill = fill_highlight if corr >= 0.4 else fill
        c16.border = thin_border

        curr_r += 1

    auto_fit_columns(ws, min_width=13)


def create_individual_user_sheet(wb, user_display, user_rows, elo_rating):
    # Sheet name max 31 chars
    safe_title = f"{user_display[:20]} Scaling"
    ws = wb.create_sheet(title=safe_title)

    # Banner
    ws.merge_cells("A1:M1")
    title_cell = ws["A1"]
    title_cell.value = f"SNAPSHOT SCALING PROFILE: @{user_display}"
    title_cell.font = font_title
    title_cell.fill = fill_title
    title_cell.alignment = align_center
    ws.row_dimensions[1].height = 32

    ws.merge_cells("A2:M2")
    sub_cell = ws["A2"]
    total_pos = user_rows[0]["total_available_positions"] if user_rows else 0
    sub_cell.value = f"Player Rating: ~{elo_rating} Elo | Total Games Extracted: 25 | Evaluated Positions: {int(total_pos):,}"
    sub_cell.font = font_subtitle
    sub_cell.fill = fill_title
    sub_cell.alignment = align_center
    ws.row_dimensions[2].height = 20

    headers = [
        "Snapshots", "Maia Top-1", "Clone Top-1", "Top-1 Δ", "Maia Top-3", "Clone Top-3",
        "Maia Loss", "Clone Loss", "Loss Δ", "Pred Tempo", "Actual Tempo", "MedAE", "Within ±2s", "Pearson r"
    ]

    ws.row_dimensions[4].height = 28
    for c_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=4, column=c_idx, value=h)
        cell.font = font_header
        cell.fill = fill_header_teal
        cell.alignment = align_header
        cell.border = thin_border

    curr_r = 5
    for r_idx, r in enumerate(user_rows):
        fill = fill_zebra if r_idx % 2 == 1 else fill_white
        ws.row_dimensions[curr_r].height = 22

        k = int(r["snapshots_evaluated"])
        is_full = k == int(total_pos)

        # Snapshots
        c1 = ws.cell(row=curr_r, column=1, value=f"{k:,} (Full)" if is_full else f"{k:,}")
        c1.alignment = align_center
        c1.font = font_bold
        c1.fill = fill
        c1.border = thin_border

        # Maia Top 1
        c2 = ws.cell(row=curr_r, column=2, value=float(r["maia_top1_pct"]) / 100.0)
        c2.number_format = "0.00%"
        c2.alignment = align_right
        c2.fill = fill
        c2.border = thin_border

        # Clone Top 1
        c3 = ws.cell(row=curr_r, column=3, value=float(r["clone_top1_pct"]) / 100.0)
        c3.number_format = "0.00%"
        c3.alignment = align_right
        c3.font = font_bold
        c3.fill = fill
        c3.border = thin_border

        # Top 1 Delta
        d1 = float(r["top1_delta_pct"]) / 100.0
        c4 = ws.cell(row=curr_r, column=4, value=d1)
        c4.number_format = "+0.00%;-0.00%;0.00%"
        c4.alignment = align_right
        c4.font = font_positive if d1 > 0 else (font_negative if d1 < 0 else font_regular)
        c4.fill = fill_highlight if d1 > 0 else fill
        c4.border = thin_border

        # Maia Top 3
        c5 = ws.cell(row=curr_r, column=5, value=float(r["maia_top3_pct"]) / 100.0)
        c5.number_format = "0.00%"
        c5.alignment = align_right
        c5.fill = fill
        c5.border = thin_border

        # Clone Top 3
        c6 = ws.cell(row=curr_r, column=6, value=float(r["clone_top3_pct"]) / 100.0)
        c6.number_format = "0.00%"
        c6.alignment = align_right
        c6.font = font_bold
        c6.fill = fill
        c6.border = thin_border

        # Maia Loss
        c7 = ws.cell(row=curr_r, column=7, value=float(r["maia_loss"]))
        c7.number_format = "0.0000"
        c7.alignment = align_right
        c7.fill = fill
        c7.border = thin_border

        # Clone Loss
        c8 = ws.cell(row=curr_r, column=8, value=float(r["clone_loss"]))
        c8.number_format = "0.0000"
        c8.alignment = align_right
        c8.font = font_bold
        c8.fill = fill
        c8.border = thin_border

        # Loss Delta
        ld = float(r["loss_reduction"])
        c9 = ws.cell(row=curr_r, column=9, value=-ld)
        c9.number_format = "-0.0000;+0.0000;0.0000"
        c9.alignment = align_right
        c9.font = font_positive if ld > 0 else (font_negative if ld < 0 else font_regular)
        c9.fill = fill_highlight if ld > 0 else fill
        c9.border = thin_border

        # Pred Tempo
        c10 = ws.cell(row=curr_r, column=10, value=float(r["pred_median_sec"]))
        c10.number_format = '0.00"s"'
        c10.alignment = align_right
        c10.fill = fill
        c10.border = thin_border

        # Actual Tempo
        c11 = ws.cell(row=curr_r, column=11, value=float(r["actual_median_sec"]))
        c11.number_format = '0.00"s"'
        c11.alignment = align_right
        c11.fill = fill
        c11.border = thin_border

        # MedAE
        c12 = ws.cell(row=curr_r, column=12, value=float(r["medae_sec"]))
        c12.number_format = '0.00"s"'
        c12.alignment = align_right
        c12.fill = fill
        c12.border = thin_border

        # Within 2s
        c13 = ws.cell(row=curr_r, column=13, value=float(r["within_2s_pct"]) / 100.0)
        c13.number_format = "0.0%"
        c13.alignment = align_right
        c13.fill = fill
        c13.border = thin_border

        # Pearson r
        corr = float(r["correlation_r"])
        c14 = ws.cell(row=curr_r, column=14, value=corr)
        c14.number_format = "0.000"
        c14.alignment = align_right
        c14.font = font_bold if corr >= 0.4 else font_regular
        c14.fill = fill_highlight if corr >= 0.4 else fill
        c14.border = thin_border

        curr_r += 1

    auto_fit_columns(ws, min_width=13)


def main():
    print(f"Reading benchmarks from {BENCHMARK_CSV}...")
    with open(BENCHMARK_CSV, "r", encoding="utf-8") as f:
        bench_rows = list(csv.DictReader(f))

    print(f"Reading scaling data from {SCALING_CSV}...")
    with open(SCALING_CSV, "r", encoding="utf-8") as f:
        scale_rows = list(csv.DictReader(f))

    wb = openpyxl.Workbook()

    # 1. Overall Benchmarks
    print("Building Sheet 1: Overall Model Benchmarks...")
    create_overall_benchmarks_sheet(wb, bench_rows)

    # 2. Master Scaling Sheet
    print("Building Sheet 2: Snapshot Scaling (All Users)...")
    create_scaling_master_sheet(wb, scale_rows)

    # 3. Individual user sheets
    user_ratings = {
        "saiteja1000": 1515,
        "Schwoby_1200": 1234,
        "hikaru": 2838,
        "stuffyness": 1788,
    }

    for user_display, rating in user_ratings.items():
        print(f"Building Individual Sheet: {user_display}...")
        u_rows = [r for r in scale_rows if r["display_username"].lower() == user_display.lower()]
        create_individual_user_sheet(wb, user_display, u_rows, rating)

    wb.save(OUTPUT_XLSX)
    print(f"[SUCCESS] Formatted Excel report generated at: {OUTPUT_XLSX}")


if __name__ == "__main__":
    main()
